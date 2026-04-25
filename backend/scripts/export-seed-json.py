"""
export-seed-json.py
Reads "STORE_FOR_WOMEN_SeedData.xlsx" (sheet: Clean Data)
and writes "seed-data.json" ready for the Node.js seed script.

File layout:
  Row 1 = English headers (CATEGORY, SUBCATEGORY, NAME, COLOR, QUANTITY, SUPPLIER, COST, SALE(WHOLESALE), SALE(RETAIL))
  Row 2 = Arabic headers (skip)
  Row 3+ = data

Run from: c:\work\magd store - v2
  python magd_store/backend/scripts/export-seed-json.py
"""

import json
from openpyxl import load_workbook
from pathlib import Path

INPUT  = Path(__file__).parent.parent.parent / "STORE_FOR_WOMEN_SeedData.xlsx"
OUTPUT = Path(__file__).parent.parent / "seed-data.json"

wb = load_workbook(INPUT, data_only=True)
ws = wb["Clean Data"]

rows = list(ws.iter_rows(values_only=True))

# Row 0 = English headers, Row 1 = Arabic headers (skip), Row 2+ = data
eng_headers = [str(h).strip().upper() if h else "" for h in rows[0]]

def ci(name):
    for i, h in enumerate(eng_headers):
        if name.upper() in h:
            return i
    raise ValueError(f"Column '{name}' not found in headers: {eng_headers}")

IDX_CATEGORY    = ci("CATEGORY")
IDX_SUBCATEGORY = ci("SUBCATEGORY")
IDX_NAME        = ci("NAME")
IDX_COLOR       = ci("COLOR")
IDX_QTY         = ci("QUANTITY")
IDX_SUPPLIER    = ci("SUPPLIER")
IDX_COST        = ci("COST")
IDX_WHOLE       = ci("WHOLESALE")
IDX_RETAIL      = ci("RETAIL")

def safe_int(v):
    if v is None:
        return 0
    try:
        return round(float(v))
    except:
        return 0

def safe_str(v):
    if v is None:
        return ""
    return str(v).strip()

records = []
for row in rows[2:]:   # skip both header rows
    category    = safe_str(row[IDX_CATEGORY])
    subcategory = safe_str(row[IDX_SUBCATEGORY])
    name        = safe_str(row[IDX_NAME])
    color       = safe_str(row[IDX_COLOR])
    qty         = safe_int(row[IDX_QTY])
    supplier    = safe_str(row[IDX_SUPPLIER])
    cost        = safe_int(row[IDX_COST])
    whole       = safe_int(row[IDX_WHOLE])
    retail      = safe_int(row[IDX_RETAIL])

    if not name or not category:
        continue

    records.append({
        "category":    category,
        "subcategory": subcategory,
        "name":        name,
        "color":       color,
        "qty":         qty,
        "supplier":    supplier,
        "cost":        cost,
        "wholesale":   whole,
        "retail":      retail,
    })

OUTPUT.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Exported {len(records)} records to {OUTPUT}")
