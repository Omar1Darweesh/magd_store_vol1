"""
clean-excel.py  -  Cleans "STORE FOR WOMEN.xlsx" before seeding.

Rules:
  Rule 1 - Exact duplicate rows -> keep first, report removed (with count)
  Rule 2 - Same (category + name + color + supplier) but different numbers
            -> flag ALL, exclude from clean, report for manual review
  Rule 3 - Same name/category but different color or supplier -> keep as-is

Seed columns (8):
  الصنف | اسم الصنف | اللون | العدد | المورد | تكلفة | جملة | قطاعي

Numbers rounded to nearest integer.

Output sheets (in order):
  1. ملخص              - summary statistics
  2. تكرارات محذوفة    - Rule 1: exact duplicate rows removed
  3. تحتاج مراجعة      - Rule 2: same key, different numbers
  4. بيانات نظيفة      - clean data ready to seed
"""

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
from collections import defaultdict

INPUT  = Path(__file__).parent.parent / "STORE FOR WOMEN.xlsx"
OUTPUT = Path(__file__).parent.parent / "STORE FOR WOMEN - CLEANED.xlsx"

# ─── Styles ─────────────────────────────────────────────────────────────────
HDR_FILL   = PatternFill("solid", fgColor="1F4E79")
HDR_FONT   = Font(bold=True, color="FFFFFF", size=11)
HDR_ALIGN  = Alignment(horizontal="center", vertical="center", wrap_text=True)

NOTE_FILL  = PatternFill("solid", fgColor="FCE4D6")   # light orange  – exact dup
PART_FILL  = PatternFill("solid", fgColor="FFF2CC")   # light yellow  – partial dup

ALT_FILL   = PatternFill("solid", fgColor="DEEAF1")   # light blue    – alt row clean

THIN = Side(style="thin", color="AAAAAA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NUM_FMT = '#,##0'   # integer with thousands separator

def style_header(ws, col_count, has_note=False):
    """Bold blue header row."""
    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.font      = HDR_FONT
        cell.fill      = HDR_FILL
        cell.alignment = HDR_ALIGN
        cell.border    = BORDER
    ws.row_dimensions[1].height = 28

def style_data_rows(ws, rows_count, col_count, fill=None, num_cols=None, has_note=False):
    """Apply borders, alternating fills, and number format."""
    num_cols = num_cols or []
    for r in range(2, rows_count + 2):
        row_fill = fill if fill else (ALT_FILL if r % 2 == 0 else None)
        for c in range(1, col_count + 1):
            cell = ws.cell(row=r, column=c)
            cell.border    = BORDER
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if row_fill:
                cell.fill = row_fill
            # Apply number format to numeric columns (skip note col=1 if present)
            real_col = c - 1 if has_note else c  # shift if note col exists
            if c in num_cols:
                cell.number_format = NUM_FMT

def auto_col_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val = str(cell.value) if cell.value is not None else ""
                max_len = max(max_len, len(val))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 8), 50)

# ─── Load source ─────────────────────────────────────────────────────────────
src_wb = load_workbook(INPUT, data_only=True)
src_ws = src_wb["store for women"]
all_rows = list(src_ws.iter_rows(values_only=True))

# Row 0 = totals, Row 1 = headers, Row 2+ = data
headers = all_rows[1]

def ci(name):
    return headers.index(name)

IDX_CATEGORY = ci("الصنف")
IDX_NAME     = ci("اسم الصنف")
IDX_COLOR    = ci("اللون")
IDX_QTY      = ci("العدد")
IDX_SUPPLIER = ci("المورد")
IDX_COST     = ci("تكلفة")
IDX_WHOLE    = ci("جملة")
IDX_RETAIL   = ci("قطاعي")

# Ordered seed columns
SEED_COLS  = [IDX_CATEGORY, IDX_NAME, IDX_COLOR, IDX_QTY,
              IDX_SUPPLIER, IDX_COST, IDX_WHOLE, IDX_RETAIL]
SEED_HEADS = [headers[i] for i in SEED_COLS]

KEY_COLS = [IDX_CATEGORY, IDX_NAME, IDX_COLOR, IDX_SUPPLIER]
NUM_COLS = [IDX_QTY, IDX_COST, IDX_WHOLE, IDX_RETAIL]

# Column positions (1-based) in output sheet WITHOUT note col
# الصنف=1, اسم الصنف=2, اللون=3, العدد=4, المورد=5, تكلفة=6, جملة=7, قطاعي=8
NUM_OUTPUT_COLS_CLEAN = [4, 6, 7, 8]       # qty, cost, wholesale, retail (no note col)
NUM_OUTPUT_COLS_NOTE  = [5, 7, 8, 9]       # same but shifted by 1 for note col

def round_row(row):
    row = list(row)
    for i in NUM_COLS:
        v = row[i]
        if isinstance(v, (int, float)) and v is not None:
            row[i] = round(v)
    return tuple(row)

def seed_key(row):
    return tuple(str(row[i]).strip() if row[i] is not None else "" for i in KEY_COLS)

def seed_numbers(row):
    return tuple(row[i] for i in NUM_COLS)

def full_fingerprint(row):
    return tuple(row[i] for i in SEED_COLS)

# Build data rows
data_rows = []
for raw in all_rows[2:]:
    if any(raw[i] is not None and str(raw[i]).strip() != "" for i in SEED_COLS):
        data_rows.append(round_row(raw))

print(f"Total rows after loading: {len(data_rows)}")

# ─── Rule 1: Exact duplicates ────────────────────────────────────────────────
seen = {}
exact_dup_rows = []
after_rule1 = []

for row in data_rows:
    fp = full_fingerprint(row)
    if fp in seen:
        seen[fp] += 1
        exact_dup_rows.append(row)
    else:
        seen[fp] = 1
        after_rule1.append(row)

total_count = dict(seen)
print(f"Rule 1 - Exact duplicates removed: {len(exact_dup_rows)}")

# ─── Rule 2: Same key, different numbers ────────────────────────────────────
key_groups = defaultdict(list)
for row in after_rule1:
    key_groups[seed_key(row)].append(row)

partial_dup_rows = []
clean_rows = []

for key, group in key_groups.items():
    if len(group) == 1:
        clean_rows.append(group[0])
    else:
        unique_nums = set(seed_numbers(r) for r in group)
        if len(unique_nums) > 1:
            partial_dup_rows.extend(group)
        else:
            clean_rows.extend(group)

print(f"Rule 2 - Partial duplicate rows flagged: {len(partial_dup_rows)}")
print(f"Clean rows ready to seed: {len(clean_rows)}")

# ─── Write workbook ──────────────────────────────────────────────────────────
out_wb = Workbook()
out_wb.remove(out_wb.active)

NOTE_HEAD = "ملاحظة"

# ── Sheet 1: Summary ─────────────────────────────────────────────────────────
ws_sum = out_wb.create_sheet("ملخص")
ws_sum.sheet_view.rightToLeft = True

summary_rows = [
    ("البيان", "العدد"),
    ("إجمالي الصفوف في الملف الأصلي",        len(data_rows)),
    ("تكرارات مطابقة تم حذفها (القاعدة 1)",   len(exact_dup_rows)),
    ("صفوف تحتاج مراجعة (القاعدة 2)",         len(partial_dup_rows)),
    ("صفوف نظيفة جاهزة للإدخال",              len(clean_rows)),
]

for r_idx, (label, val) in enumerate(summary_rows, start=1):
    ws_sum.cell(row=r_idx, column=1, value=label)
    ws_sum.cell(row=r_idx, column=2, value=val)

# Header style
for c in range(1, 3):
    cell = ws_sum.cell(row=1, column=c)
    cell.font      = HDR_FONT
    cell.fill      = HDR_FILL
    cell.alignment = HDR_ALIGN
    cell.border    = BORDER

# Data rows style
fills = [None, NOTE_FILL, PART_FILL, PatternFill("solid", fgColor="E2EFDA")]
for r in range(2, len(summary_rows) + 1):
    f = fills[r - 2] if (r - 2) < len(fills) else None
    for c in range(1, 3):
        cell = ws_sum.cell(row=r, column=c)
        cell.border    = BORDER
        cell.alignment = Alignment(horizontal="right", vertical="center")
        if f:
            cell.fill = f
        if c == 2 and r > 1:
            cell.number_format = NUM_FMT
            cell.font = Font(bold=True, size=11)

ws_sum.column_dimensions["A"].width = 42
ws_sum.column_dimensions["B"].width = 16
ws_sum.row_dimensions[1].height = 28

# ── Helper: write a data sheet ───────────────────────────────────────────────
def write_data_sheet(ws, rows, notes=None, row_fill=None):
    ws.sheet_view.rightToLeft = True
    has_note = notes is not None
    col_count = len(SEED_HEADS) + (1 if has_note else 0)
    num_cols_set = set(NUM_OUTPUT_COLS_NOTE if has_note else NUM_OUTPUT_COLS_CLEAN)

    # ── Header row ────────────────────────────────────────────────────────────
    header = ([NOTE_HEAD] + list(SEED_HEADS)) if has_note else list(SEED_HEADS)
    for c, val in enumerate(header, start=1):
        cell = ws.cell(row=1, column=c, value=val)
        cell.font      = HDR_FONT
        cell.fill      = HDR_FILL
        cell.alignment = HDR_ALIGN
        cell.border    = BORDER
    ws.row_dimensions[1].height = 28

    # ── Data rows – write value AND style in one pass ─────────────────────────
    for idx, row in enumerate(rows):
        r = idx + 2
        values = [row[i] for i in SEED_COLS]
        if has_note:
            values = [notes[idx]] + values

        alt = PatternFill("solid", fgColor="DEEAF1") if r % 2 == 0 else None

        for c, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border    = BORDER
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if row_fill:
                cell.fill = row_fill
            elif alt:
                cell.fill = alt
            if c in num_cols_set:
                cell.number_format = NUM_FMT

    auto_col_width(ws)


# ── Sheet 2: Exact duplicates ─────────────────────────────────────────────────
ws_dup = out_wb.create_sheet("تكرارات محذوفة")
exact_notes = [
    f"ظهر {total_count.get(full_fingerprint(r), 2)} مرات - احتُفظ بنسخة واحدة"
    for r in exact_dup_rows
] if exact_dup_rows else []

write_data_sheet(ws_dup,
                 exact_dup_rows if exact_dup_rows else [],
                 notes=exact_notes if exact_dup_rows else None,
                 row_fill=NOTE_FILL if exact_dup_rows else None)

if not exact_dup_rows:
    ws_dup.append(["لا توجد تكرارات مطابقة"])

# ── Sheet 3: Partial duplicates ───────────────────────────────────────────────
ws_part = out_wb.create_sheet("تحتاج مراجعة")
partial_notes = [
    "نفس الصنف/اللون/المورد - أرقام مختلفة"
    for _ in partial_dup_rows
] if partial_dup_rows else []

write_data_sheet(ws_part,
                 partial_dup_rows if partial_dup_rows else [],
                 notes=partial_notes if partial_dup_rows else None,
                 row_fill=PART_FILL if partial_dup_rows else None)

if not partial_dup_rows:
    ws_part.append(["لا توجد تكرارات جزئية"])

# ── Sheet 4: Clean seed data ──────────────────────────────────────────────────
ws_clean = out_wb.create_sheet("بيانات نظيفة")
write_data_sheet(ws_clean, clean_rows)

# ─── Save ─────────────────────────────────────────────────────────────────────
out_wb.save(OUTPUT)

print(f"\nDone! Output: {OUTPUT}")
print(f"\nSummary:")
print(f"  Total input rows            : {len(data_rows)}")
print(f"  Rule 1 (exact dup) removed  : {len(exact_dup_rows)}")
print(f"  Rule 2 (partial dup) flagged: {len(partial_dup_rows)}")
print(f"  Clean rows to seed          : {len(clean_rows)}")
print(f"\nSeed columns: {', '.join(SEED_HEADS)}")
